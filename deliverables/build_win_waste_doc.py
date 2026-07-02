from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "deliverables/win_waste_material_intelligence_cloud.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(92, 101, 116)
GRAY_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_border_bottom(paragraph, color="2E74B5", size="12", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    r.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def add_para(doc, text="", style=None, before=None, after=None, align=None):
    p = doc.add_paragraph(text, style=style)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_callout(doc, label, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [9120])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label)
    set_run_font(r, size=11, color=INK, bold=True)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    add_para(doc, "", after=4)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [2100, 7140])
    hdr = table.rows[0].cells
    hdr[0].text = "Dimension"
    hdr[1].text = "Recommendation"
    for cell in hdr:
        set_cell_shading(cell, GRAY_FILL)
        for p in cell.paragraphs:
            for r in p.runs:
                set_run_font(r, size=10.5, bold=True, color=INK)
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text = k
        cells[1].text = v
        for p in cells[0].paragraphs:
            for r in p.runs:
                set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
        for p in cells[1].paragraphs:
            for r in p.runs:
                set_run_font(r, size=10.5, color=INK)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    add_para(doc, "", after=4)


def add_comparison_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    for idx, h in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = h
        set_cell_shading(cell, GRAY_FILL)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                set_run_font(r, size=10, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = text
            for p in cells[idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (2, 3) else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_run_font(r, size=9.5, color=INK)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=90, bottom=90)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    add_para(doc, "", after=4)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.10

    for list_style in ("List Bullet", "List Number"):
        st = styles[list_style]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.167

    hdr = section.header.paragraphs[0]
    hdr.text = "WIN Waste Innovations | Application Opportunity Analysis"
    hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in hdr.runs:
        set_run_font(r, size=9, color=MUTED)
    ftr = section.footer.paragraphs[0]
    ftr.text = "Confidential strategy concept | Prepared June 9, 2026"
    ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in ftr.runs:
        set_run_font(r, size=9, color=MUTED)


def title_page(doc):
    add_para(doc, "STRATEGY BRIEF", before=18, after=4)
    title = add_para(doc, "WIN Material Intelligence Cloud", after=2)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.runs[0]
    set_run_font(r, size=24, color=INK, bold=True)
    subtitle = add_para(
        doc,
        "A high-value application concept for WIN Waste Innovations that converts operations data into customer retention, premium revenue, ESG proof, and routing intelligence.",
        after=12,
    )
    for r in subtitle.runs:
        set_run_font(r, size=13, color=MUTED)
    rows = [
        ("Prepared for", "Exploratory product and strategy discussion"),
        ("Prepared date", "June 9, 2026"),
        ("Primary recommendation", "Build an integrated waste-stream intelligence platform around WIN's curb-to-grid, rail, recycling, special waste, and WTE capabilities."),
        ("Document type", "Executive product opportunity brief"),
    ]
    for label, value in rows:
        p = add_para(doc, after=2)
        lr = p.add_run(f"{label}: ")
        set_run_font(lr, size=10.5, color=INK, bold=True)
        vr = p.add_run(value)
        set_run_font(vr, size=10.5, color=INK)
    rule = add_para(doc, "", before=10, after=10)
    set_paragraph_border_bottom(rule)
    add_callout(
        doc,
        "Recommendation in one sentence",
        "WIN should productize its operational and sustainability data into a live customer and operations platform that tells every account where its material went, what value it created, how service can be optimized, and what verified ESG claims the customer can safely use.",
    )


def build_doc():
    doc = Document()
    configure_document(doc)
    title_page(doc)

    add_heading(doc, "1. Executive Summary", 1)
    add_para(
        doc,
        "WIN Waste Innovations' public positioning is unusually strong: the company is not only a waste hauler, but an integrated collection, recycling, waste-to-energy, rail logistics, landfill, special waste, and sustainability operator. The most valuable application opportunity is therefore not another generic waste portal. It is a material intelligence layer that makes WIN's integrated infrastructure visible, measurable, and monetizable for customers and internal teams.",
    )
    add_para(
        doc,
        "The proposed application, WIN Material Intelligence Cloud, would create a verified digital account of each customer's waste stream. It would show material volumes, recycling outcomes, WTE conversion, recovered metals, rail vs. truck movement, landfill diversion, cost-saving recommendations, contamination issues, and ESG-ready outputs. For WIN, this becomes a customer-retention engine, premium analytics product, sales differentiator, and operational optimizer.",
    )
    add_key_value_table(
        doc,
        [
            ("Strategic thesis", "WIN already owns the physical network. The next value layer is verified, customer-facing material intelligence."),
            ("Ideal first users", "Commercial portfolios, municipalities, universities, hospitals, manufacturers, food processors, and high-volume compactor accounts."),
            ("Business model", "Bundled enterprise portal for strategic accounts plus premium analytics/reporting tiers for commercial and municipal customers."),
            ("Primary impact", "Higher retention, premium revenue, fewer unnecessary hauls, better diversion, stronger sales conversion, and defensible sustainability reporting."),
        ],
    )

    add_heading(doc, "2. What the Public Site Reveals", 1)
    add_para(
        doc,
        "The website presents several clear product and operations signals. These signals point toward a platform that integrates customer experience, operations optimization, and sustainability proof.",
    )
    add_bullets(
        doc,
        [
            "WIN promotes residential, business, compactor, roll-off, recycling, special waste, and industry-specific services.",
            "The Green by Design and curb-to-grid messaging says WIN optimizes pickup routes, diverts waste from landfills, recovers metals, converts remaining waste into renewable power, and provides customers with a detailed summary of sustainability contributions.",
            "The company highlights waste-by-rail as cost- and fuel-efficient, with a large private railcar fleet and millions of tons moved annually.",
            "The sustainability report publishes strong aggregate metrics, including tons converted to renewable energy, tons moved by low-carbon rail, tons recycled, metals recovered, and customers served.",
            "The compactor business page emphasizes fewer hauls, lower costs, safety, onsite evaluation, and waste audits.",
            "The special-waste page emphasizes compliant handling, end-to-end logistics, destruction certificates, flexible scheduling, and special handling of complex non-hazardous waste streams.",
        ],
    )

    add_heading(doc, "3. Strategic Gap", 1)
    add_para(
        doc,
        "The public story is compelling, but the visible digital experience appears to stop short of making that story operationally interactive for customers. The site has contact channels, account access, quote flows, and a sustainability calculator, but the high-value opportunity is a persistent customer-specific operating record: what happened to my material, how did it create value, what can I improve next, and what evidence can I export?",
    )
    add_callout(
        doc,
        "Core gap",
        "WIN's differentiator is integrated infrastructure. The product opportunity is to expose that infrastructure as verified, account-level intelligence that customers and internal teams can act on every month.",
    )

    add_heading(doc, "4. Proposed Application", 1)
    add_heading(doc, "Product Name: WIN Material Intelligence Cloud", 2)
    add_para(
        doc,
        "WIN Material Intelligence Cloud would be both a customer portal and an internal decision-support platform. Customers would receive a live sustainability and service ledger. WIN teams would receive operational recommendations for routing, destination selection, compactor sizing, contamination intervention, rail utilization, special-waste workflow management, and customer advisory opportunities.",
    )
    add_heading(doc, "Primary User Jobs", 2)
    add_bullets(
        doc,
        [
            "Customer sustainability lead: export verified waste, diversion, recycling, WTE, and renewable-energy contribution data for ESG, procurement, board, LEED, municipal, or tenant reporting.",
            "Facilities or operations manager: reduce overflows, missed pickups, unnecessary hauls, contamination events, and service uncertainty.",
            "WIN sales team: demonstrate quantifiable savings and sustainability value during proposals and renewals.",
            "WIN dispatcher or logistics planner: choose better pickup timing and destination paths based on fill level, route economics, rail availability, facility capacity, and customer commitments.",
            "WIN account manager: identify upsell opportunities such as compactors, Green Glove/WTE paths, cardboard diversion, special-waste handling, and reporting packages.",
        ],
    )

    add_heading(doc, "5. Core Modules", 1)
    add_comparison_table(
        doc,
        ["Module", "What it does", "Primary users", "Value created"],
        [
            ("Customer Sustainability Ledger", "Shows collected tons, recycled materials, WTE conversion, metals recovered, rail movement, landfill diversion, and monthly/annual exports.", "Customers, account managers", "Retention, ESG proof, premium reporting revenue"),
            ("Smart Container and Compactor Optimization", "Uses pickup history, scale tickets, fill telemetry, seasonality, and service issues to recommend container mix and pickup cadence.", "Facilities teams, dispatch, sales", "Fewer unnecessary hauls, lower cost, higher compactor conversion"),
            ("Waste Stream Diagnostics", "Uses photos, intake observations, and facility data to identify contamination, recoverable commodities, prohibited materials, and special-waste opportunities.", "Customers, MRF/transfer station teams", "Higher diversion, lower contamination cost, safer handling"),
            ("Curb-to-Grid Destination Optimizer", "Recommends destination paths across recycling, transfer, rail, WTE, landfill, or special handling based on cost, capacity, compliance, and sustainability impact.", "Dispatch, operations, logistics", "Better margin, capacity use, rail utilization, and customer promise fulfillment"),
            ("Premium Advisory Layer", "Turns waste audits into recurring recommendations, benchmarks, certificates, compliance artifacts, and what-if simulations.", "Strategic accounts, sales", "Paid advisory tier, stronger renewals, differentiated proposals"),
        ],
        [2200, 3720, 1500, 1940],
    )

    add_heading(doc, "6. MVP Definition", 1)
    add_para(
        doc,
        "The MVP should avoid boiling the ocean. Start with one dense region and one high-value segment: commercial customers using compactors, roll-offs, or recurring business waste/recycling service. Massachusetts, Connecticut, or New York would be natural candidates because WIN has a concentration of hauling, transfer, WTE, and rail assets in the broader Northeast.",
    )
    add_heading(doc, "MVP Capabilities", 2)
    add_numbered(
        doc,
        [
            "Account dashboard: customer sites, containers, service schedule, recent pickups, weights, issues, and invoice/service summary.",
            "Sustainability summary: tons collected, estimated material pathway, recycling, WTE, rail movement, landfill diversion, and renewable power equivalents.",
            "Export pack: monthly PDF and CSV outputs for ESG, facilities reporting, LEED support, procurement, and municipal summaries.",
            "Container intelligence: recommended pickup frequency, container changes, compactor candidates, overflow alerts, and underutilized-service flags.",
            "Photo-assisted diagnostics: customer or WIN employee uploads images of waste areas, contamination, overflow, special waste, or recoverable materials.",
            "Internal operations view: map account material flows to transfer stations, rail-served facilities, WTE, recycling, landfill, and special-waste paths.",
            "What-if simulator: model cost and sustainability changes from moving to compactor service, adding cardboard diversion, adjusting pickup cadence, or routing via rail/WTE.",
        ],
    )

    add_heading(doc, "7. Data Foundation", 1)
    add_para(
        doc,
        "The application should start with existing data and add sensors selectively. WIN does not need every customer container instrumented on day one. The value begins by unifying data already created by operations.",
    )
    add_comparison_table(
        doc,
        ["Data source", "Examples", "Initial use", "Maturity"],
        [
            ("Customer/account data", "Customer, site, service type, contract, container inventory", "Account views and segmentation", "MVP"),
            ("Route and service data", "Scheduled stops, completed stops, missed stops, service exceptions", "Reliability, route analysis, customer alerts", "MVP"),
            ("Scale and ticket data", "Transfer station weights, WTE intake, landfill tickets, MRF outputs", "Material ledger and destination proof", "MVP"),
            ("Rail logistics data", "Railcar movement, transfer station origin, destination, tonnage", "Rail impact and cost/carbon modeling", "MVP+"),
            ("Compactor telemetry", "Fill level, pressure, cycles, fault signals", "Pickup optimization and maintenance", "Pilot"),
            ("Photo/computer vision", "Dock images, contaminated bins, special-waste intake, MRF snapshots", "Diagnostics and customer coaching", "Pilot"),
            ("Sustainability factors", "WTE equivalents, rail vs. truck factors, recycling factors, landfill methane assumptions", "ESG calculations and reporting", "MVP"),
        ],
        [1900, 2820, 2920, 1720],
    )

    add_heading(doc, "8. Reference Architecture", 1)
    add_para(
        doc,
        "A practical architecture should layer on top of existing systems rather than replacing them. The platform should ingest from operational systems, normalize material events into a shared ledger, then expose role-specific applications.",
    )
    add_bullets(
        doc,
        [
            "Ingestion: CRM, billing, route management, scale systems, transfer station systems, WTE facility records, MRF outputs, rail logistics, customer service records, and optional IoT devices.",
            "Material event ledger: a canonical record of collection, transfer, processing, recovery, rail movement, WTE conversion, landfill disposal, exception, and certificate events.",
            "Analytics layer: service optimization, compactor recommendations, contamination scoring, customer benchmarking, sustainability calculations, and what-if modeling.",
            "Applications: customer portal, account-manager cockpit, dispatch/operations dashboard, sales demo mode, ESG export generator, and special-waste certificate workflow.",
            "Governance: audit trails, factor versioning, customer-visible assumptions, approval workflows, and role-based access.",
        ],
    )

    add_heading(doc, "9. Business Value Model", 1)
    add_para(
        doc,
        "The value case is strongest because the same platform can create revenue, margin, retention, and brand value. It does not depend on one narrow ROI lever.",
    )
    add_comparison_table(
        doc,
        ["Value lever", "How the app creates value", "Example metric"],
        [
            ("Premium revenue", "Sell analytics, reporting, advisory, and verified ESG exports as subscription tiers.", "Attach rate and monthly recurring revenue per strategic account"),
            ("Retention", "Make WIN the source of record for waste and sustainability data.", "Churn reduction among enrolled customers"),
            ("Haul optimization", "Avoid underfilled pickups and dispatch only when service value is clear.", "Hauls avoided and cost saved"),
            ("Compactor conversion", "Identify customers where compactors reduce cost and improve service.", "New compactor deals and avoided overflow events"),
            ("Rail utilization", "Show and optimize rail movement where feasible.", "Tons moved by rail and truck miles avoided"),
            ("Contamination reduction", "Detect problem streams early and coach customers.", "Contamination events and rejected loads reduced"),
            ("Sales conversion", "Use customer-specific simulations in proposals.", "Win rate and average contract value"),
        ],
        [1800, 5100, 2460],
    )

    add_heading(doc, "10. Rollout Roadmap", 1)
    add_comparison_table(
        doc,
        ["Phase", "Duration", "Scope", "Exit criteria"],
        [
            ("Discovery and data audit", "4-6 weeks", "Map source systems, customer segments, high-value workflows, sustainability factors, and reporting needs.", "Data inventory, MVP region, pilot accounts, and calculation methodology selected"),
            ("MVP build", "10-14 weeks", "Customer dashboard, event ledger, exports, service recommendations, and internal account view.", "20-50 pilot accounts live with monthly reports"),
            ("Pilot expansion", "8-12 weeks", "Add compactor telemetry, photo diagnostics, sales demo mode, and account-manager workflows.", "Measured reduction in reporting effort and at least one operational savings lever proven"),
            ("Scale product", "6-9 months", "Broaden regions, add rail optimization, special-waste workflows, APIs, and premium tiers.", "Repeatable revenue packaging and operational adoption across regions"),
        ],
        [1500, 1500, 3900, 2460],
    )

    add_heading(doc, "11. Risk Register", 1)
    add_comparison_table(
        doc,
        ["Risk", "Why it matters", "Mitigation"],
        [
            ("Calculation trust", "Customers may rely on reports for ESG or procurement claims.", "Version calculation factors, show assumptions, keep audit trails, and distinguish measured vs. estimated values."),
            ("Data fragmentation", "Route, scale, facility, rail, and billing systems may not align cleanly.", "Start with a canonical material event ledger and pilot only where source data quality is sufficient."),
            ("Operational adoption", "Dispatch and facility teams may resist another dashboard.", "Make recommendations advisory at first and integrate into existing workflows instead of forcing a new daily system."),
            ("Sensor economics", "IoT everywhere can be expensive and slow.", "Instrument only high-volume compactors and problematic sites first; use existing tickets for the first MVP."),
            ("Customer sensitivity", "Some accounts may not want contamination or disposal problems exposed.", "Frame diagnostics as savings and improvement opportunities; allow account-manager review before customer release."),
            ("Regulatory claims", "Sustainability outputs can create reputational risk if overstated.", "Use conservative language, source every factor, and route formal claims through approved templates."),
        ],
        [1900, 3500, 3960],
    )

    add_heading(doc, "12. Recommended Next Step", 1)
    add_para(
        doc,
        "Run a focused 30-day product discovery sprint with one regional operating team and 10-15 representative commercial accounts. The sprint should inventory available data, reconstruct each customer's prior three months of material flow, mock the dashboard and ESG export, and quantify two operational levers: haul optimization and compactor/right-sizing opportunity.",
    )
    add_callout(
        doc,
        "Decision point",
        "If the sprint can produce customer-ready reports from existing data for at least 70% of pilot accounts, proceed to MVP. If not, prioritize the material event ledger and source-system cleanup before adding advanced AI or IoT.",
    )

    add_heading(doc, "13. Source Notes", 1)
    sources = [
        ("WIN Waste homepage", "https://www.win-waste.com/"),
        ("WIN Waste About Us", "https://www.win-waste.com/about-us/"),
        ("WIN Waste Green by Design", "https://www.win-waste.com/green-by-design/"),
        ("WIN Waste Waste-by-Rail", "https://www.win-waste.com/green-by-design/waste-by-rail/"),
        ("WIN Waste Sustainability Report page", "https://www.win-waste.com/community-and-responsibility/sustainability-report/"),
        ("WIN Waste Business Services", "https://www.win-waste.com/business/"),
        ("WIN Waste Commercial Compactor Rental", "https://www.win-waste.com/business/compactors/"),
        ("WIN Waste Waste Disposal / Special Waste Services", "https://www.win-waste.com/business/special-waste-service/"),
        ("WIN Waste Locations", "https://www.win-waste.com/about-us/locations/"),
        ("U.S. EIA: Waste-to-energy", "https://www.eia.gov/energyexplained/biomass/waste-to-energy.php"),
        ("U.S. EPA: Basic Information about Landfill Gas", "https://www.epa.gov/lmop/basic-information-about-landfill-gas"),
        ("U.S. EPA: Municipal Solid Waste Landfills NSPS/EG", "https://www.epa.gov/stationary-sources-air-pollution/municipal-solid-waste-landfills-new-source-performance-standards"),
        ("ScienceDirect: AI and IoT in solid waste management review", "https://www.sciencedirect.com/science/article/pii/S2590123025001069"),
    ]
    for name, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{name}: ")
        hyperlink(p, url, url)

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
