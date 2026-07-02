from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "deliverables/win_waste_tanstack_implementation_plan.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(92, 101, 116)
GREEN = RGBColor(29, 96, 76)
RED = RGBColor(155, 28, 28)
GRAY_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
GREEN_FILL = "EAF4EF"
CALLOUT_FILL = "F4F6F9"
RISK_FILL = "FFF1F0"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def paragraph_border_bottom(paragraph, color="2E74B5", size="12", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    r.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def add_para(doc, text="", style=None, before=None, after=None, align=None):
    p = doc.add_paragraph(text, style=style)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_callout(doc, label, body, fill=CALLOUT_FILL, color=INK):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [9120])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label)
    set_run_font(r, size=11, color=color, bold=True)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    for r in p2.runs:
        set_run_font(r, size=10.5, color=INK)
    add_para(doc, "", after=3)


def add_table(doc, headers, rows, widths, center_cols=None, header_fill=GRAY_FILL):
    center_cols = center_cols or []
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, header_fill)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                set_run_font(r, size=9.5, color=INK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_run_font(r, size=9.2, color=INK)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    add_para(doc, "", after=4)
    return table


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.10

    for list_style in ("List Bullet", "List Number"):
        st = styles[list_style]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.167

    hdr = section.header.paragraphs[0]
    hdr.text = "WIN Material Intelligence Cloud | TanStack Implementation Plan"
    hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in hdr.runs:
        set_run_font(r, size=9, color=MUTED)
    ftr = section.footer.paragraphs[0]
    ftr.text = "World-class execution blueprint | Prepared June 9, 2026"
    ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in ftr.runs:
        set_run_font(r, size=9, color=MUTED)


def cover(doc):
    add_para(doc, "IMPLEMENTATION BLUEPRINT", before=18, after=4)
    title = add_para(doc, "WIN Material Intelligence Cloud")
    title.paragraph_format.space_after = Pt(2)
    set_run_font(title.runs[0], size=24, color=INK, bold=True)
    subtitle = add_para(
        doc,
        "A fool-proof, phase-gated plan to build a strategic waste-stream intelligence platform using TanStack Start, Router, Query, DB, Table, Virtual, Form, and Store.",
        after=12,
    )
    for r in subtitle.runs:
        set_run_font(r, size=13, color=MUTED)
    meta = [
        ("Audience", "WIN executive sponsors, product leadership, engineering, data, operations, sales, account management, and compliance teams"),
        ("Technical posture", "TanStack-first application layer with deliberate backend boundaries, version pinning, auditability, and operational fail-safes"),
        ("Planning standard", "Phase-gated implementation with entry criteria, exit criteria, acceptance tests, rollback paths, and executive decision points"),
        ("Prepared date", "June 9, 2026"),
    ]
    for label, value in meta:
        p = add_para(doc, after=2)
        lr = p.add_run(f"{label}: ")
        set_run_font(lr, size=10.5, color=INK, bold=True)
        vr = p.add_run(value)
        set_run_font(vr, size=10.5, color=INK)
    rule = add_para(doc, "", before=10, after=10)
    paragraph_border_bottom(rule)
    add_callout(
        doc,
        "The promise",
        "This plan makes the product hard to derail: it starts with business proof, creates a canonical material event ledger before advanced AI, uses TanStack where it is strongest, and requires measurable gates before scaling.",
        fill=GREEN_FILL,
        color=GREEN,
    )


def build():
    doc = Document()
    configure(doc)
    cover(doc)

    doc.add_heading("1. Executive Implementation Doctrine", 1)
    add_para(
        doc,
        "The goal is to build a live system of record for customer waste streams, not a dashboard theater project. The implementation must prove business value before expanding technical ambition. Every phase therefore has measurable customer, operations, data, security, and engineering gates.",
    )
    add_bullets(
        doc,
        [
            "Start with the material event ledger. If material movements cannot be trusted, every dashboard, AI model, ESG report, and optimization recommendation becomes fragile.",
            "Use TanStack for the application contract: typed routes, URL state, server/client data boundaries, cache behavior, high-volume tables, validated forms, and reactive client collections.",
            "Keep sustainability calculations auditable. Every customer-visible claim must carry a calculation version, source event IDs, assumptions, measured vs. estimated classification, and approval status.",
            "Treat operations adoption as a product feature. Dispatchers, account managers, facility teams, and customer sustainability leads need different workflows, not one universal screen.",
            "Pilot with a deliberately narrow operating surface and expand only after the platform beats manual reporting, identifies savings, and survives real account-manager review.",
        ],
    )

    doc.add_heading("2. Recommended TanStack Architecture", 1)
    add_para(
        doc,
        "TanStack should own the frontend and full-stack application experience while backend data pipelines, calculation services, identity integration, and operational systems remain explicitly bounded. This avoids turning the web framework into the enterprise integration layer.",
    )
    add_table(
        doc,
        ["TanStack capability", "Role in the platform", "Implementation decision", "Why it matters"],
        [
            ("TanStack Start", "Full-stack application shell with SSR, streaming, server functions, and server routes.", "Use for portal, internal cockpit, server-side loaders, and typed server actions; pin versions because Start is RC.", "Gives a unified app model without surrendering client-side interactivity."),
            ("TanStack Router", "Typed route tree, params, route loaders, authenticated routes, and URL search state.", "Make URL state the source for filters, tabs, date ranges, account IDs, and report views.", "Deep links become reliable artifacts for account reviews and operations handoffs."),
            ("TanStack Query", "Server-state cache, background refetching, optimistic mutations, invalidation, and hydration.", "Use for APIs, report status, service events, recommendations, and customer records.", "Prevents manual cache drift across large dashboards."),
            ("TanStack DB", "Reactive client collections and live queries over API-backed data.", "Use selectively for account workspaces, material events, recommendations, and offline-friendly review queues.", "Enables spreadsheet-like responsiveness without inventing a local query engine."),
            ("TanStack Table", "Headless engine for enterprise grids.", "Use for material event tables, exceptions, account portfolios, invoices, facilities, and recommendation queues.", "Provides controlled sorting, filtering, grouping, column visibility, selection, and pagination."),
            ("TanStack Virtual", "Virtualization for huge lists and grids.", "Pair with Table for high-volume material events, route stops, and facility intake logs.", "Keeps UI fast when users inspect thousands of rows."),
            ("TanStack Form", "Validated forms and workflow submissions.", "Use for waste audits, special-waste intake, report approvals, customer settings, and field photo diagnostics.", "Keeps complex forms type-safe and validation-rich."),
            ("TanStack Store", "Small client-side UI/workflow state.", "Use only for transient UI state that is not server data or URL state.", "Avoids misusing Query as a general client-state dump."),
        ],
        [1800, 2400, 2900, 2260],
    )

    doc.add_heading("3. Platform Blueprint", 1)
    add_para(
        doc,
        "The architecture should have five durable layers. Each layer has a contract and can be tested independently.",
    )
    add_numbers(
        doc,
        [
            "Source-system integration layer: route systems, scale tickets, CRM, billing, WTE intake, MRF output, rail logistics, landfill tickets, customer service records, compactor telemetry, photo uploads, and special-waste certificates.",
            "Canonical material event ledger: immutable append-only events for collection, transfer, sorting, processing, recycling, rail movement, WTE conversion, landfill disposal, recovered metal, exception, certificate, and calculation events.",
            "Calculation and recommendation services: sustainability factors, routing scenarios, compactor right-sizing, contamination diagnostics, ESG exports, what-if modeling, and recommendation scoring.",
            "Application APIs and server functions: typed server boundaries used by TanStack Start routes, with access control, audit logging, idempotency, and report generation.",
            "Role-specific product surfaces: customer portal, account-manager cockpit, dispatch dashboard, facility exception queue, sales demo mode, executive analytics, and compliance review console.",
        ],
    )

    doc.add_heading("4. Fool-Proof Phase Gates", 1)
    add_table(
        doc,
        ["Phase", "Name", "Primary objective", "Exit gate"],
        [
            ("0", "Mandate and Operating Model", "Create sponsorship, scope, product principles, and decision rights.", "Named executive sponsor, product owner, technical owner, data owner, compliance owner, and pilot region."),
            ("1", "Discovery and Proof Inventory", "Map real workflows, systems, data quality, customer needs, and reporting obligations.", "Pilot accounts selected; data availability scored; manual baseline measured."),
            ("2", "Data Foundation", "Build material event ledger and ingestion contracts before UI scale-up.", "Three months of pilot events reconciled with invoices/tickets above agreed accuracy threshold."),
            ("3", "TanStack Product Shell", "Create the secure app foundation, route model, design system, auth, and observability.", "Authenticated route tree, role guards, test harness, deployment pipeline, and baseline pages live."),
            ("4", "MVP Intelligence Modules", "Deliver customer ledger, exports, service recommendations, tables, and account cockpit.", "Pilot customers and account managers can complete monthly review without spreadsheets."),
            ("5", "Optimization and AI Assist", "Add diagnostics, compactor intelligence, what-if modeling, and recommendation governance.", "Recommendations show measurable savings or diversion lift and are human-approved."),
            ("6", "Pilot Hardening", "Run real operations, resolve support loops, tighten calculations, and prove commercial model.", "Executive go/no-go based on adoption, value, accuracy, security, and margin signals."),
            ("7", "Regional Scale", "Expand to more accounts, regions, facilities, and premium packages.", "Repeatable onboarding playbook and SLA-backed support model."),
            ("8", "Enterprise Flywheel", "Turn the platform into a durable competitive moat.", "Platform becomes a renewal, sales, advisory, and operations operating system."),
        ],
        [720, 1950, 4200, 2490],
        center_cols=[0],
        header_fill=BLUE_FILL,
    )

    phases = [
        (
            "Phase 0: Mandate and Operating Model",
            "2 weeks",
            "Prevent ambiguity before engineering starts.",
            [
                "Define the strategic product charter: customer value, operational value, revenue hypothesis, and explicit non-goals.",
                "Name DRIs for product, engineering, data, operations, compliance, cybersecurity, sales, and customer success.",
                "Select one pilot region, one customer segment, and one service mix; avoid national scope at inception.",
                "Create a decision forum that can resolve data ownership, calculation policy, customer claims, and operational workflow disputes within 48 hours.",
            ],
            [
                "Signed product charter.",
                "Pilot region and 20-50 candidate accounts.",
                "Calculation governance owner.",
                "Approved technology posture, including TanStack Start version pinning and fallback plan.",
            ],
        ),
        (
            "Phase 1: Discovery and Proof Inventory",
            "4-6 weeks",
            "Find the exact places where data can become money, retention, or operational savings.",
            [
                "Shadow account managers, dispatchers, facility teams, sales reps, customer service, and sustainability/reporting staff.",
                "Trace five real customer material journeys from pickup through transfer, rail, WTE, recycling, landfill, invoice, and report.",
                "Inventory source systems, identifiers, timestamps, reconciliation keys, latency, owner, and known trust gaps.",
                "Baseline manual reporting time, current service exceptions, missed/overflow events, contamination issues, and compactor opportunities.",
                "Prototype reports in static form and validate with pilot customers before building dynamic screens.",
            ],
            [
                "Data map with red/yellow/green confidence by source.",
                "Top 10 workflows ranked by value and feasibility.",
                "Report mockups approved by customers and account managers.",
                "MVP backlog frozen for Phase 2 and Phase 3.",
            ],
        ),
        (
            "Phase 2: Data Foundation and Material Event Ledger",
            "8-12 weeks",
            "Create the trust layer that the entire product will stand on.",
            [
                "Define canonical event types, event IDs, source IDs, account/site/container/facility references, timestamps, units, evidence links, and correction semantics.",
                "Build ingestion adapters for CRM/account data, route/service events, scale tickets, WTE intake, MRF output, landfill tickets, rail movement, and billing.",
                "Create reconciliation jobs that compare material events against invoices, facility totals, and route completion records.",
                "Add data-quality scoring, exception queues, source lineage, duplicate detection, and late-arriving event handling.",
                "Define sustainability calculation versions and separate measured values from estimated values.",
            ],
            [
                "Pilot accounts have at least 90 days of reconstructed material events.",
                "Reconciliation accuracy meets the executive-approved threshold.",
                "Every customer-visible metric can trace back to source events and calculation versions.",
                "Data defects flow to an owner, not a dead-letter pile.",
            ],
        ),
        (
            "Phase 3: TanStack Product Shell",
            "6-8 weeks, parallel with late Phase 2",
            "Build the secure, testable application platform before feature sprawl.",
            [
                "Scaffold TanStack Start with React, TypeScript, strict mode, Biome or ESLint, route generation, CI, and pinned package versions.",
                "Model routes around real workspaces: /accounts, /accounts/$accountId/sites/$siteId, /ledger, /reports, /recommendations, /dispatch, /facilities, /admin.",
                "Use TanStack Router search params for date range, account filters, facility filters, tab state, table state, and saved views.",
                "Use TanStack Query for server-state calls and define query-key factories before feature teams create ad hoc keys.",
                "Establish TanStack Table primitives, Virtual list patterns, Form wrappers, accessibility checks, error boundaries, and role-based route guards.",
                "Instrument performance, errors, slow queries, mutation failure rates, and user journeys from day one.",
            ],
            [
                "App shell deploys to staging with authentication, roles, observability, and automated tests.",
                "Route contracts and query keys are documented and linted.",
                "A representative table handles 50,000 row-equivalent data with virtualization and usable filtering.",
                "No customer data is visible without enforced role and tenant checks.",
            ],
        ),
        (
            "Phase 4: MVP Intelligence Modules",
            "12-16 weeks",
            "Deliver enough end-to-end product to replace spreadsheet-heavy monthly review.",
            [
                "Customer Sustainability Ledger: account-level tons, service events, material pathways, recycling, WTE, rail movement, landfill diversion, and exportable summaries.",
                "Account Manager Cockpit: account health, exceptions, missing data, renewal talking points, recommendations, and report approval workflow.",
                "Service Optimization View: pickup cadence, missed service, overflow, underfilled haul signals, compactor candidacy, and container right-sizing recommendations.",
                "Report Export Engine: PDF/CSV outputs with calculation version, source period, approval status, and customer-visible assumptions.",
                "Operations Exception Queue: reconcile failures, missing tickets, suspicious weights, duplicate events, late events, and customer-facing correction workflow.",
            ],
            [
                "Pilot account managers can run monthly reviews without manual spreadsheet assembly.",
                "Customers can self-serve approved reports and understand assumptions.",
                "Data exceptions are visible, triaged, and resolved within agreed SLA.",
                "MVP proves at least two quantified value levers, such as report-time reduction and haul-optimization savings.",
            ],
        ),
        (
            "Phase 5: Optimization and AI Assist",
            "10-14 weeks",
            "Add intelligence only where the ledger and workflows have already earned trust.",
            [
                "Introduce compactor telemetry and fill-level optimization for high-value accounts.",
                "Add photo-assisted diagnostics for contamination, overflow, special-waste evidence, and recoverable material opportunities.",
                "Build what-if simulations for pickup frequency, compactor conversion, cardboard diversion, WTE path, rail movement, and service changes.",
                "Create human-in-the-loop recommendation approval with confidence, explanation, data lineage, and expected value.",
                "Add sales demo mode that can safely simulate outcomes without exposing unapproved claims.",
            ],
            [
                "Recommendations have measured precision/recall or business acceptance scores.",
                "No recommendation becomes customer-visible without approval rules.",
                "At least one optimization creates statistically credible savings or service improvement.",
                "Model outputs include explanation, confidence, and rollback path.",
            ],
        ),
        (
            "Phase 6: Pilot Hardening and Commercial Packaging",
            "8-10 weeks",
            "Turn a working MVP into a product WIN can sell, support, and defend.",
            [
                "Run weekly pilot account reviews and capture adoption, confusion, missing metrics, report edits, support tickets, and value outcomes.",
                "Finalize pricing tiers: included portal, premium ESG exports, advisory package, API access, and enterprise portfolio analytics.",
                "Run security review, privacy review, claim-language review, and disaster recovery tabletop exercise.",
                "Create onboarding, training, support scripts, sales collateral, QBR templates, and customer success playbooks.",
            ],
            [
                "Pilot customers renew or express willingness to pay for premium value.",
                "Support model handles real questions without engineering intervention.",
                "Security and compliance sign off on data exposure, claims, and auditability.",
                "Executive team approves regional scale.",
            ],
        ),
        (
            "Phase 7: Regional Scale",
            "6-9 months",
            "Scale the product without letting data quality, support, or custom reporting chaos outrun the platform.",
            [
                "Create a repeatable region-onboarding playbook: source systems, data mapping, facility profiles, calculation factors, account segmentation, and training.",
                "Expand from pilot accounts to strategic commercial portfolios and municipalities.",
                "Add API and scheduled delivery options for enterprise ESG systems and customer procurement portals.",
                "Introduce premium benchmark views across sites, regions, industries, and service types.",
                "Track product-qualified expansion signals for compactors, recycling, special waste, Green Glove/WTE service, and advisory work.",
            ],
            [
                "New region onboarding time falls below target threshold.",
                "Portal usage and report downloads correlate with account retention or expansion.",
                "Support load per customer decreases as self-service maturity improves.",
                "Premium attach rate validates the revenue model.",
            ],
        ),
        (
            "Phase 8: Enterprise Flywheel",
            "Ongoing",
            "Convert the product from a useful portal into a strategic moat.",
            [
                "Use aggregate learning to improve routing, rail utilization, facility capacity planning, contamination coaching, and commodity recovery.",
                "Create executive portfolio intelligence for customers with multi-site operations.",
                "Offer verified APIs and data feeds to customer ESG, procurement, finance, and facilities systems.",
                "Build a marketplace of expert playbooks: hospital waste, university move-out, manufacturing scrap, food processor diversion, municipal reporting, and commercial real estate tenant reporting.",
            ],
            [
                "Platform influences renewals, pricing, routing decisions, and capital planning.",
                "WIN can show differentiated proof that competitors cannot easily reproduce.",
                "The material event ledger becomes a durable internal data asset.",
                "Operational insights feed back into service design and infrastructure utilization.",
            ],
        ),
    ]

    doc.add_heading("5. Detailed Phase Plan", 1)
    for title, duration, objective, stages, exits in phases:
        doc.add_heading(title, 2)
        add_para(doc, f"Duration: {duration}", after=2)
        add_para(doc, f"Objective: {objective}", after=6)
        doc.add_heading("Stages", 3)
        add_bullets(doc, stages)
        doc.add_heading("Exit Criteria", 3)
        add_bullets(doc, exits)

    doc.add_heading("6. Application Information Architecture", 1)
    add_table(
        doc,
        ["Area", "Routes", "Primary TanStack patterns", "Acceptance standard"],
        [
            ("Customer Portal", "/accounts/$accountId/overview, /reports, /ledger, /recommendations", "Router params/search, Query hydration, Table, Form", "Customer can understand month-to-date performance and download an approved report in under 3 minutes."),
            ("Account Manager Cockpit", "/internal/accounts, /internal/accounts/$accountId/review", "Router authenticated routes, Query mutations, Table selection, DB collections", "Account manager can review exceptions, approve report, and capture customer follow-up without spreadsheets."),
            ("Operations Dashboard", "/internal/dispatch, /internal/facilities, /internal/exceptions", "Virtualized tables, controlled filters, live queries", "Teams can triage high-priority exceptions without full-page reloads or stale state."),
            ("Recommendation Center", "/recommendations, /recommendations/$id", "Query, Form approval, Store for review workflow", "Every recommendation has expected value, confidence, source events, approval state, and rollback path."),
            ("Admin and Governance", "/admin/calculations, /admin/roles, /admin/sources", "Protected routes, validated forms, audit tables", "Calculation factors, role grants, and source mappings are versioned and auditable."),
        ],
        [1700, 2850, 2600, 2210],
    )

    doc.add_heading("7. Engineering Standards", 1)
    add_bullets(
        doc,
        [
            "TypeScript strict mode is mandatory. No untyped payload crosses the API boundary.",
            "All route params and search params are validated. Invalid URL state degrades gracefully.",
            "Every query key uses a centralized factory. Cache invalidation is reviewed like database schema design.",
            "Every mutation has idempotency, optimistic UI rules, error recovery, and audit logging where customer-visible state changes.",
            "Every table has an explicit empty state, loading state, error state, pagination/virtualization strategy, and saved-view behavior.",
            "Every reportable metric has source lineage, calculation version, measured/estimated flag, and owner.",
            "Every feature has instrumentation before launch: latency, error rate, adoption, completion rate, and business-value events.",
            "Accessibility, keyboard navigation, and responsive behavior are release gates, not polish tasks.",
        ],
    )

    doc.add_heading("8. Test and Verification Matrix", 1)
    add_table(
        doc,
        ["Layer", "Tests", "Release-blocking examples"],
        [
            ("Data ingestion", "Schema validation, duplicate detection, replay tests, reconciliation, late-event tests", "Scale ticket missing account mapping; duplicate route events; event totals do not reconcile."),
            ("Ledger and calculations", "Golden datasets, factor-version tests, measured vs. estimated tests, audit-trail tests", "Customer report cannot trace a metric to source events or factor version."),
            ("TanStack app", "Route tests, search-param tests, query-cache tests, mutation rollback tests, table state tests", "User sees wrong account data after navigation; stale report appears after approval."),
            ("Security", "Tenant isolation, role tests, object authorization, audit logs, penetration testing", "Customer A can infer Customer B data through URL, cache, export, or API."),
            ("UX and workflow", "Task completion, support scripts, pilot user testing, accessibility checks", "Account manager needs spreadsheet workaround for approved MVP workflow."),
            ("Operations", "Load tests, failover, backup restore, observability drills, incident runbooks", "Report jobs silently fail; event backlog grows without alerting."),
        ],
        [1700, 4400, 3260],
    )

    doc.add_heading("9. Staffing Model", 1)
    add_table(
        doc,
        ["Role", "Phase intensity", "Responsibilities"],
        [
            ("Executive sponsor", "High in Phases 0, 6, 7", "Set mandate, approve scope, unblock organizational decisions, own commercial stakes."),
            ("Product lead", "High throughout", "Own roadmap, user research, prioritization, acceptance, launch plan, and value metrics."),
            ("Engineering lead", "High throughout", "Own architecture, delivery model, code quality, reliability, and technical tradeoffs."),
            ("Data architect", "Highest in Phases 1-3", "Design ledger, source contracts, reconciliation, lineage, and data-quality strategy."),
            ("Frontend/TanStack engineers", "Highest in Phases 3-5", "Build routes, tables, forms, query patterns, app shell, and role-specific workflows."),
            ("Backend/platform engineers", "Highest in Phases 2-6", "Build ingestion, APIs, server functions, queues, reports, auth integration, and observability."),
            ("Operations SMEs", "High in Phases 1, 4, 6", "Validate workflows, exception handling, dispatch/facility needs, and adoption risks."),
            ("Compliance/security", "Review gates each phase", "Approve claims, data exposure, auditability, privacy, and access control."),
            ("Customer success/sales", "High in Phases 4-8", "Pilot customers, packaging, demos, renewals, adoption, and commercial proof."),
        ],
        [2200, 1900, 5260],
    )

    doc.add_heading("10. Risk Controls", 1)
    add_table(
        doc,
        ["Risk", "Early warning", "Control"],
        [
            ("TanStack Start maturity", "Framework churn, plugin changes, migration friction", "Pin versions, isolate framework-specific code, maintain a Router/Query-compatible fallback, and upgrade only at phase boundaries."),
            ("Data trust failure", "Manual corrections dominate reports", "Block customer expansion until ledger reconciliation and source ownership are stable."),
            ("Dashboard without behavior change", "Users view pages but keep old spreadsheets", "Define task completion as the adoption metric, not page visits."),
            ("Overpromised sustainability claims", "Marketing language outruns calculation certainty", "Use approved claim templates and measured/estimated labels."),
            ("AI before data maturity", "Model demos look impressive but fail real reviews", "AI cannot enter customer workflow until event lineage and human approval are proven."),
            ("Enterprise integration drag", "Legacy source systems slow delivery", "Use thin source adapters and reconcile in the ledger instead of redesigning source systems."),
            ("Support explosion", "Every customer wants custom report edits", "Create template governance, premium advisory tier, and strict report-variant rules."),
        ],
        [2300, 2600, 4460],
        header_fill=RISK_FILL,
    )

    doc.add_heading("11. Metrics That Decide Whether This Wins", 1)
    add_table(
        doc,
        ["Metric", "Target interpretation", "Owner"],
        [
            ("Manual reporting hours avoided", "Core MVP value; should drop materially for pilot accounts.", "Product + account management"),
            ("Ledger reconciliation confidence", "Trust gate; should be high before customer-scale reporting.", "Data owner"),
            ("Customer report downloads and shares", "Proof that customers treat the platform as useful evidence.", "Customer success"),
            ("Recommendations accepted", "Proof that analytics create operational action.", "Operations + product"),
            ("Hauls avoided or right-sized", "Direct margin and customer-cost lever.", "Operations"),
            ("Premium attach rate", "Commercial proof for analytics/reporting tier.", "Sales"),
            ("Churn delta for enrolled accounts", "Strategic retention proof.", "Executive sponsor"),
            ("Support tickets per active account", "Scale readiness signal.", "Customer success"),
            ("P95 route/page latency", "App-quality signal for heavy operational workflows.", "Engineering"),
            ("Security/access incidents", "Must remain zero for cross-customer exposure.", "Security"),
        ],
        [3000, 4200, 2160],
    )

    doc.add_heading("12. Practical First 30 Days", 1)
    add_numbers(
        doc,
        [
            "Name the executive sponsor, product lead, engineering lead, data owner, compliance reviewer, and pilot operations lead.",
            "Choose one pilot region and 20-50 candidate accounts with compactors, recurring service, recycling, and enough ticket history.",
            "Pull three months of route, scale, invoice, transfer/WTE/MRF/landfill, and account data for 10 representative accounts.",
            "Create the first canonical material event model and reconcile it against invoices and facility totals.",
            "Design static customer report mockups and review them with account managers and two friendly customers.",
            "Scaffold the TanStack Start app shell, route tree, auth stub, query-key conventions, table primitives, and CI pipeline.",
            "Hold the first gate review: continue only if the data can plausibly produce trusted customer reports.",
        ],
    )

    doc.add_heading("13. Source Notes", 1)
    sources = [
        ("TanStack Start official documentation", "https://tanstack.com/start/latest"),
        ("TanStack Start routing guide", "https://tanstack.com/start/latest/docs/framework/react/guide/routing"),
        ("TanStack Router official documentation", "https://tanstack.com/router/latest/docs"),
        ("TanStack Router and Query integration", "https://tanstack.com/router/latest/docs/integrations/query"),
        ("TanStack DB quick start", "https://tanstack.com/db/latest/docs/quick-start"),
        ("TanStack Table overview", "https://tanstack.com/table/beta/docs/overview"),
        ("TanStack Virtual documentation", "https://tanstack.com/virtual/latest/docs"),
        ("TanStack Form validation guide", "https://tanstack.com/form/v1/docs/framework/react/guides/validation"),
        ("WIN Waste homepage", "https://www.win-waste.com/"),
        ("WIN Waste Green by Design", "https://www.win-waste.com/green-by-design/"),
        ("WIN Waste Sustainability Report", "https://www.win-waste.com/community-and-responsibility/sustainability-report/"),
    ]
    for name, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{name}: ")
        hyperlink(p, url, url)

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
